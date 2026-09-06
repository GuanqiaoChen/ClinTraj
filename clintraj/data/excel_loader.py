"""Faithful, read-only Excel ingestion with safe errors and explicit provenance.

This module does not infer problem ownership, correct graph labels, apply reviewer
edits, or expose full records as agent observations. The source is a reference
corpus; clinical validity and actual availability times remain unestablished.
"""

import hashlib
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal, overload

import openpyxl
from pydantic import BaseModel, ConfigDict, ValidationError

from clintraj.data.review import ClinicianReview, DecisionRating, RationaleRating
from clintraj.domain.action_types import ActionType
from clintraj.domain.decision_graph import ClinicalDecisionGraph
from clintraj.domain.relation_types import RelationType
from clintraj.domain.schemas import GoldenEdge, GoldenNode

MAIN_SHEET = "待标注数据"
SOURCE_SHA256 = "23e067a5de05a5fe119a541cf90fcc2f40399b5b538131058f4a939497a2741d"
GOLDEN_CASE_HASHES = (
    "5b6e050e83a4b6c025aec7dec0b8f752a5b590c6a12c6d84dca1edf1cfbbdf0a",
    "e58643b89e1f64c313d6408dc900481839b300fd374b7883ecceb71f3718fa1d",
    "50bcdce028b0502d7068618ddfedb74aa35d43d6beb4970d8e9bc5a8a0d8598c",
    "898ef30e3401cad09142e6b89de4ce27de0e9845c08180cba61cd3c4d5979428",
    "e0dc8fce1628b3f2746d0b0f8ea22a1c96e860944c3d73d595a94deaa1bf4ad6",
)
CANONICAL_HEADERS = (
    "病案标识",
    "步骤id",
    "父步骤id",
    "路径关系",
    "动作类型",
    "上一个决策动作后解锁的新信息",
    "决策（给医生的建议动作）",
    "临床理由",
)
REVIEW_HEADERS = (
    "医生_决策正确性(正确填1，正确但有其他更推荐决策填2，不正确填3)",
    "医生_修订决策",
    "医生_临床理由质量()",
    "医生_修订临床理由",
)
INSTRUCTION_MARKERS = (
    "系统字段",
    "系统字段",
    "系统字段",
    "系统字段；START/CONTINUE/BRANCH/CONSULT/TRANSFER/RETURN",
    "系统字段",
)


class DataValidationError(ValueError):
    """Safe diagnostic: never includes cell values, narratives, or raw case IDs."""

    def __init__(self, code: str, *, row: int | None = None) -> None:
        self.code = code
        self.row = row
        super().__init__(f"{code}" + (f" at Excel row {row}" if row is not None else ""))


class IngestionEvent(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)
    code: str
    source_row: int | None = None
    case_ref: str | None = None
    step_id: int | None = None


class DataInspectionReport(BaseModel):
    """Aggregate, shareable report; excludes clinical narrative and source IDs."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    workbook_sha256: str
    sheet_name: str
    sheet_names: tuple[str, ...]
    case_count: int
    node_count: int
    reviewed_node_count: int
    annotated_node_count: int
    action_counts: dict[str, int]
    relation_counts: dict[str, int]
    events: tuple[IngestionEvent, ...] = ()
    source_review_status: str = "review_labels_absent_and_optional_for_technical_execution"
    temporal_basis: str = "recorded_step_order_without_verified_release_timestamps"


@dataclass(frozen=True)
class LoadedDataset:
    graphs: tuple[ClinicalDecisionGraph, ...] = field(repr=False)
    reviews: tuple[ClinicianReview, ...] = field(repr=False)
    report: DataInspectionReport

    def case(self, case_id: str) -> ClinicalDecisionGraph:
        for graph in self.graphs:
            if graph.case_id == case_id:
                return graph
        raise KeyError("case_not_found")

    def case_by_hash(self, case_sha256: str) -> ClinicalDecisionGraph:
        """Resolve a local source case without embedding its identifier in public fixtures."""
        if len(case_sha256) != 64:
            raise KeyError("case_not_found")
        for graph in self.graphs:
            if hashlib.sha256(graph.case_id.encode("utf-8")).hexdigest() == case_sha256:
                return graph
        raise KeyError("case_not_found")


@overload
def _integer(value: Any, *, row: int, nullable: Literal[False] = False) -> int: ...


@overload
def _integer(value: Any, *, row: int, nullable: Literal[True]) -> int | None: ...


def _integer(value: Any, *, row: int, nullable: bool = False) -> int | None:
    if value is None and nullable:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise DataValidationError("expected_numeric_integer", row=row)
    if not isinstance(value, int) and not value.is_integer():
        raise DataValidationError("expected_numeric_integer", row=row)
    if value < 1:
        raise DataValidationError("expected_positive_integer", row=row)
    return int(value)


@overload
def _text(value: Any, *, row: int, optional: Literal[False] = False) -> str: ...


@overload
def _text(value: Any, *, row: int, optional: Literal[True]) -> str | None: ...


def _text(value: Any, *, row: int, optional: bool = False) -> str | None:
    if value is None and optional:
        return None
    if not isinstance(value, str) or not value.strip():
        raise DataValidationError("expected_nonempty_text", row=row)
    return value


def parse_record(
    values: tuple[Any, ...], *, row_number: int
) -> tuple[GoldenNode, GoldenEdge, ClinicianReview]:
    """Validate one fixed-layout source record without narrative-changing repair."""
    if len(values) != 12:
        raise DataValidationError("expected_twelve_columns", row=row_number)
    case_id = _text(values[0], row=row_number)
    step_id = _integer(values[1], row=row_number)
    parent = _integer(values[2], row=row_number, nullable=True)
    try:
        relation = RelationType(values[3])
        action_type = ActionType(values[4])
    except (ValueError, TypeError):
        raise DataValidationError("unknown_ontology_label", row=row_number) from None
    evidence = _text(values[5], row=row_number)
    action = _text(values[6], row=row_number)
    rationale = _text(values[7], row=row_number)
    decision = _integer(values[8], row=row_number, nullable=True)
    revised_action = _text(values[9], row=row_number, optional=True)
    rationale_rating = _integer(values[10], row=row_number, nullable=True)
    revised_rationale = _text(values[11], row=row_number, optional=True)
    try:
        node = GoldenNode(
            case_id=case_id,
            step_id=step_id,
            new_evidence=evidence,
            action_type=action_type,
            action=action,
            clinical_rationale=rationale,
        )
        edge = GoldenEdge(parent_step_id=parent, child_step_id=step_id, relation=relation)
        review = ClinicianReview(
            case_id=case_id,
            step_id=step_id,
            source_row=row_number,
            decision_rating=DecisionRating(decision) if decision is not None else None,
            revised_action=revised_action,
            rationale_rating=RationaleRating(rationale_rating) if rationale_rating is not None else None,
            revised_rationale=revised_rationale,
        )
    except (ValueError, ValidationError):
        raise DataValidationError("invalid_canonical_record_or_review", row=row_number) from None
    return node, edge, review


def _semantic_events(graph: ClinicalDecisionGraph, *, case_ref: str) -> list[IngestionEvent]:
    """Flag review needs; do not confuse ordinal graph validity with semantics."""
    events: list[IngestionEvent] = []
    expected_parent_action = {
        RelationType.CONSULT: ActionType.CONSULT,
        RelationType.TRANSFER: ActionType.TRANSFER,
    }
    for edge in graph.edges:
        if edge.relation in expected_parent_action and edge.parent_step_id is not None:
            if graph.node(edge.parent_step_id).action_type != expected_parent_action[edge.relation]:
                events.append(IngestionEvent(
                    code="specialty_edge_requires_parent_action_adjudication",
                    case_ref=case_ref,
                    step_id=edge.child_step_id,
                ))
    if hashlib.sha256(graph.case_id.encode("utf-8")).hexdigest() == GOLDEN_CASE_HASHES[1] and not any(
        edge.relation == RelationType.RETURN for edge in graph.edges
    ):
        events.append(IngestionEvent(code="selected_case_2_missing_image_return", case_ref=case_ref))
    return events


def load_workbook(path: str | Path, *, sheet_name: str = MAIN_SHEET) -> LoadedDataset:
    """Load the clinical source once, preserve every source edge, close read-only.

    Sheet selection is explicit. The teaching-example sheet is not inferred to
    be a trajectory cohort. External links/macros are neither followed nor run.
    """
    source = Path(path)
    try:
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=False, keep_links=False)
    except Exception:
        raise DataValidationError("unreadable_workbook") from None

    nodes: dict[str, list[GoldenNode]] = defaultdict(list)
    edges: dict[str, list[GoldenEdge]] = defaultdict(list)
    reviews: list[ClinicianReview] = []
    events: list[IngestionEvent] = []
    seen: set[tuple[str, int]] = set()
    first_rows: dict[str, int] = {}
    sheet_names = tuple(workbook.sheetnames)
    try:
        if sheet_name not in workbook.sheetnames:
            raise DataValidationError("required_sheet_missing")
        worksheet = workbook[sheet_name]
        records = worksheet.iter_rows()
        header_cells = next(records, ())
        header = tuple(cell.value for cell in header_cells)
        if header != CANONICAL_HEADERS + REVIEW_HEADERS:
            raise DataValidationError("unexpected_header_schema", row=1)
        for row_number, cells in enumerate(records, start=2):
            if any(cell.data_type in ("f", "e") for cell in cells):
                raise DataValidationError("formula_or_error_cell_not_allowed", row=row_number)
            values = tuple(cell.value for cell in cells)
            if not any(value is not None for value in values):
                events.append(IngestionEvent(code="skip_empty_row", source_row=row_number))
                continue
            if row_number == 2 and values[:5] == INSTRUCTION_MARKERS:
                events.append(IngestionEvent(code="skip_explicit_instruction_row", source_row=2))
                continue
            node, edge, review = parse_record(values, row_number=row_number)
            key = (node.case_id, node.step_id)
            if key in seen:
                raise DataValidationError("duplicate_case_step", row=row_number)
            seen.add(key)
            first_rows.setdefault(node.case_id, row_number)
            nodes[node.case_id].append(node)
            edges[node.case_id].append(edge)
            reviews.append(review)
        if not nodes:
            raise DataValidationError("no_trajectory_records")
        graphs = []
        for case_number, (case_id, case_nodes) in enumerate(nodes.items(), start=1):
            try:
                graph = ClinicalDecisionGraph(nodes=tuple(case_nodes), edges=tuple(edges[case_id]))
                graph.validate_integrity()
            except (ValueError, ValidationError):
                raise DataValidationError("invalid_case_graph", row=first_rows[case_id]) from None
            graphs.append(graph)
            events.extend(_semantic_events(graph, case_ref=f"case_{case_number:04d}"))
    finally:
        workbook.close()

    action_counts = Counter(node.action_type.value for graph in graphs for node in graph.nodes)
    relation_counts = Counter(edge.relation.value for graph in graphs for edge in graph.edges)
    report = DataInspectionReport(
        workbook_sha256=digest,
        sheet_name=sheet_name,
        sheet_names=sheet_names,
        case_count=len(graphs),
        node_count=len(reviews),
        reviewed_node_count=sum(review.is_complete for review in reviews),
        annotated_node_count=sum(review.has_annotation for review in reviews),
        action_counts=dict(sorted(action_counts.items())),
        relation_counts=dict(sorted(relation_counts.items())),
        events=tuple(events),
    )
    return LoadedDataset(graphs=tuple(graphs), reviews=tuple(reviews), report=report)
